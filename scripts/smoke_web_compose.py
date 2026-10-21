"""Real Compose/WAL smoke, isolated project and volumes; no real shops or Telegram.

Run from the repository root after installing .[dev,web]. Requires Docker Compose.
The generated token is never printed. The isolated stack is always removed.
"""
import io
import json
import os
import secrets
import subprocess
import time
import urllib.error
import urllib.request

from openpyxl import load_workbook


def main():
    token = secrets.token_urlsafe(32)
    env = dict(os.environ, PRICE_MONITOR_WEB_TOKEN=token, PRICE_MONITOR_WEB_PORT="18000",
               PRICE_MONITOR_INTERVAL_SECONDS="3", PRICE_MONITOR_ENABLE_XCOM="0",
               PRICE_MONITOR_ENABLE_ALLO="0", PRICE_MONITOR_ENABLE_BOOKS="0")
    project = f"price-monitor-web-smoke-{os.getpid()}"
    command = ["docker", "compose", "-p", project]
    def compose(*args, **kwargs):
        return subprocess.run(command + list(args), env=env, check=True, **kwargs)
    def get(path, authenticated=True):
        request = urllib.request.Request("http://127.0.0.1:18000" + path,
                                        headers={"Authorization": f"Bearer {token}"} if authenticated else {})
        return urllib.request.urlopen(request, timeout=10)
    def wait_for(check):
        deadline = time.monotonic() + 90
        while time.monotonic() < deadline:
            try:
                value = check()
                if value: return value
            except (OSError, ValueError):
                pass
            time.sleep(.5)
        raise AssertionError("Timed out waiting for dashboard / scheduled collection")
    try:
        compose("up", "--build", "-d")
        wait_for(lambda: json.load(get("/api/history/demo/00123"))["observations"])
        try:
            get("/", authenticated=False)
            raise AssertionError("Unauthenticated dashboard was accessible")
        except urllib.error.HTTPError as exc:
            assert exc.code == 401
        for path in ["/", "/sources/demo", "/products/demo/00123", "/events", "/report", "/status"]:
            assert get(path).status == 200
        # Read continuously between real scheduled runs while monitor is alive.
        # This catches a :ro mount when the writer closes and -shm/-wal disappear.
        for _ in range(30):
            assert get("/sources/demo").status == 200
            time.sleep(.2)
        compose("exec", "-T", "demo", "price-monitor", "demo-scenario", "price_drop")
        points = wait_for(lambda: (p if any(x["price"] == "19.90" for x in p) else None)
                          if (p := json.load(get("/api/history/demo/00123"))["observations"]) else None)
        assert points[0]["price"] == "24.90"
        assert b"PRICE_CHANGED" in get("/events").read()
        response = get("/report/download?source=demo&since=2000-01-01&until=2100-01-01")
        book = load_workbook(io.BytesIO(response.read()))
        current = {r[1]: r for r in book["Current"].iter_rows(min_row=2, values_only=True)}
        assert current["00123"][4] == 19.9 and current["123"][4] == 3.5
        assert book["History"].max_row >= 3
        book.close()
        compose("exec", "-T", "web", "python", "-c", "from price_monitor.web.db import read_database; from pathlib import Path; import sqlite3;\nwith read_database(Path('/data/price_monitor.db')) as c:\n c.execute('PRAGMA query_only=OFF')\n try: c.execute('DELETE FROM observations')\n except sqlite3.OperationalError as e: assert 'readonly' in str(e)\n else: raise AssertionError('Web wrote to database')")
        print("PASS: live Compose monitor, repeated WAL reads, auth, six pages, price_drop, XLSX, mode=ro write refusal")
    finally:
        subprocess.run(command + ["logs", "--tail", "60", "monitor", "web"], env=env, check=False)
        subprocess.run(command + ["down", "-v", "--remove-orphans"], env=env, check=False)


if __name__ == "__main__":
    main()
