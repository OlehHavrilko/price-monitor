from decimal import Decimal

import httpx
import pytest
from fastapi.testclient import TestClient

from conftest import make_source
from price_monitor import demo_site, storage
from price_monitor.collector import collect_once
from price_monitor.web.app import create_app

TOKEN = "test-only-web-token"


@pytest.fixture
def web_data(tmp_path):
    path = tmp_path / "web.db"
    source = make_source("http://demo.test")
    state = tmp_path / "scenario.json"
    writer = storage.connect(path)

    def collect(scenario):
        demo_site.write_scenario(state, scenario)
        def handler(request):
            page = int(request.url.params.get("page", "1"))
            body = demo_site.render_page(page, demo_site.read_scenario(state))
            return httpx.Response(200 if body else 404, text=body or "")
        result = collect_once(writer, source, httpx.MockTransport(handler))
        assert result.ok, result.error
        return result

    collect("initial")
    storage.set_rule(writer, "demo", "00123", threshold=Decimal("20"), drop_percent=Decimal("10"))
    collect("price_drop")
    app = create_app(db_path=path, token=TOKEN, sources={"demo": source})
    with TestClient(app, headers={"Authorization": f"Bearer {TOKEN}"}) as client:
        yield client, writer, collect, path
    writer.close()


@pytest.mark.parametrize("url", ["/", "/sources/demo", "/products/demo/00123", "/events", "/report", "/status"])
def test_pages_render_demo(web_data, url):
    client, _, _, _ = web_data
    response = client.get(url)
    assert response.status_code == 200
    assert '<html lang="uk">' in response.text
    assert "Price Monitor" in response.text


def test_history_and_string_keys(web_data):
    client, _, _, _ = web_data
    lamp = client.get("/api/history/demo/00123").json()
    bulb = client.get("/api/history/demo/123").json()
    assert lamp["key"] == "00123" and bulb["key"] == "123"
    assert [r["price"] for r in lamp["observations"]] == ["24.90", "19.90"]
    assert [r["price"] for r in bulb["observations"]] == ["3.50", "3.50"]
    assert "19.90" in client.get("/products/demo/00123").text
    assert "3.50" in client.get("/products/demo/123").text


@pytest.mark.parametrize("url", ["/sources/missing", "/products/demo/missing", "/api/history/demo/missing", "/missing"])
def test_not_found(web_data, url):
    assert web_data[0].get(url).status_code == 404


def test_authentication_cookie_and_header(web_data):
    client = web_data[0]
    client.headers.pop("Authorization")
    assert client.get("/").status_code == 401
    assert client.get("/login").status_code == 200
    assert client.post("/login", data={"token": "wrong"}).status_code == 401
    response = client.post("/login", data={"token": TOKEN}, follow_redirects=False)
    assert response.status_code == 303
    assert "HttpOnly" in response.headers["set-cookie"]
    assert "SameSite=strict" in response.headers["set-cookie"]
    assert client.get("/").status_code == 200
    assert client.get("/", headers={"Authorization": "Bearer wrong"}).status_code == 401
    assert client.post("/logout").status_code == 200
    assert client.get("/").status_code == 401


def test_filter_and_partial(web_data):
    client = web_data[0]
    response = client.get("/sources/demo?q=00123", headers={"HX-Request": "true"})
    assert response.status_code == 200 and "<html" not in response.text
    assert "Desk lamp Nord" in response.text and "spare bulb" not in response.text
    assert "Office chair Basic" in client.get("/sources/demo?availability=OUT_OF_STOCK").text
    assert "Desk lamp Nord" not in client.get("/sources/demo?availability=OUT_OF_STOCK").text


def test_report_download_reuses_generator_and_cleans_files(web_data, monkeypatch):
    from io import BytesIO
    from openpyxl import load_workbook
    from price_monitor import reports
    client = web_data[0]
    generated = []
    original = reports.write_report
    def record(*args, **kwargs):
        path = original(*args, **kwargs)
        generated.append(path)
        return path
    monkeypatch.setattr(reports, "write_report", record)
    response = client.get("/report/download", params={"source": "demo", "since": "2000-01-01", "until": "2100-01-01"})
    assert response.status_code == 200
    assert generated and not generated[0].parent.exists()
    book = load_workbook(BytesIO(response.content))
    assert book.sheetnames == ["Current", "History"]
    current = {r[1]: r for r in book["Current"].iter_rows(min_row=2, values_only=True)}
    assert current["00123"][4] == 19.9 and current["123"][4] == 3.5
    assert len(list(book["History"].iter_rows(min_row=2))) == 4
    book.close()


@pytest.mark.parametrize("since,until", [("bad", "2026-01-01"), ("2026-02-01", "2026-01-01")])
def test_report_invalid_period(web_data, since, until):
    assert web_data[0].get("/report/download", params={"source":"demo", "since":since, "until":until}).status_code == 400
