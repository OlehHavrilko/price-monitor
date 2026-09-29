from datetime import datetime, timedelta, timezone
from decimal import Decimal as D

import pytest
from openpyxl import load_workbook

from conftest import make_source
from price_monitor import demo_site, reports, storage
from price_monitor.collector import collect_once
from price_monitor.models import Availability, ProductObservation

WINDOW = (datetime.now(timezone.utc) - timedelta(days=1), datetime.now(timezone.utc) + timedelta(days=1))


def test_report_matches_database_and_keeps_keys_as_text(demo, conn, tmp_path):
    url, state = demo
    collect_once(conn, make_source(url))
    demo_site.write_scenario(state, "price_drop")
    collect_once(conn, make_source(url))

    path = reports.write_report(conn, tmp_path / "out", "demo", *WINDOW)
    wb = load_workbook(path)
    current = {row[1]: row for row in wb["Current"].iter_rows(min_row=2, values_only=True)}
    assert len(current) == 12
    assert "00123" in current and "123" in current  # text keys, leading zeros kept
    assert current["00123"][4] == pytest.approx(19.90)
    assert wb["Current"]["E2"].number_format == "0.00"
    assert isinstance(wb["Current"]["H2"].value, datetime)
    assert [c.value for c in wb["History"][1]][:2] == ["Source", "Key"]

    history = [r for r in wb["History"].iter_rows(min_row=2, values_only=True)]
    assert {(r[1], r[8], r[9], r[10]) for r in history} == {
        ("00123", "PRICE_CHANGED", "24.90", "19.90"),
        ("01003", "PRICE_CHANGED", "12.50", "11.00"),
    }


def test_text_from_pages_never_becomes_a_formula(conn, tmp_path):
    storage.ensure_source(conn, "demo", "http://x.test")
    run_id = storage.start_run(conn, "demo")
    items = [
        ProductObservation("=1+1", '=HYPERLINK("http://evil.test","click")', "http://x.test/a", D("1"), "USD", Availability.IN_STOCK),
        ProductObservation("@sum", "+cmd|calc", "http://x.test/b", D("2"), "USD", Availability.UNKNOWN),
        ProductObservation("-5", "- minus", "http://x.test/c", D("3"), "USD", Availability.OUT_OF_STOCK),
    ]
    storage.commit_snapshot(conn, run_id, "demo", items)

    path = reports.write_report(conn, tmp_path, "demo", *WINDOW)
    ws = load_workbook(path)["Current"]
    for row in ws.iter_rows(min_row=2):
        for cell in row[:4]:
            assert cell.data_type != "f", f"{cell.coordinate} became a formula"
    assert {ws.cell(r, 2).value for r in range(2, 5)} == {"=1+1", "@sum", "-5"}
    names = {ws.cell(r, 2).value: ws.cell(r, 3).value for r in range(2, 5)}
    assert names["=1+1"].startswith("=HYPERLINK")


def test_empty_report_has_headers_and_no_rows(conn, tmp_path):
    path = reports.write_report(conn, tmp_path, "demo", *WINDOW)
    wb = load_workbook(path)
    for name, width in (("Current", 8), ("History", 11)):
        rows = list(wb[name].iter_rows(values_only=True))
        assert len(rows) == 1 and len(rows[0]) == width


def test_period_must_be_valid_and_report_size_is_limited(demo, conn, tmp_path, monkeypatch):
    start, end = WINDOW
    with pytest.raises(reports.ReportError):
        reports.write_report(conn, tmp_path, "demo", end, start)
    with pytest.raises(reports.ReportError):
        reports.write_report(conn, tmp_path, "demo", start.replace(tzinfo=None), end)
    collect_once(conn, make_source(demo[0]))
    monkeypatch.setattr(reports, "MAX_REPORT_ROWS", 5)
    with pytest.raises(reports.ReportError, match="limit"):
        reports.write_report(conn, tmp_path, "demo", start, end)
    assert not list(tmp_path.glob("*.xlsx"))
