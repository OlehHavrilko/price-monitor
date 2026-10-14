import sqlite3
from datetime import datetime, timedelta, timezone
from decimal import Decimal as D

from openpyxl import load_workbook

from price_monitor import reports, storage
from price_monitor.comparison import PRICE_CHANGED, Rule, compare
from price_monitor.models import Availability as A
from price_monitor.models import ProductObservation


def obs(price, availability=A.IN_STOCK):
    return ProductObservation("p1", "Band", "http://x.test/p1", None if price is None else D(price), "UAH", availability)


def types(old, new, rule=None):
    events, _ = compare({"p1": old}, {"p1": new}, {"p1"}, {}, {"p1": rule} if rule else {})
    return [(e.type, e.old_value, e.new_value) for e in events]


def test_price_appearing_or_disappearing_is_history_only_and_never_triggers_rules():
    rule = Rule(threshold=D("100"), drop_percent=D("10"))
    assert types(obs(None), obs("50"), rule) == [(PRICE_CHANGED, None, "50")]
    assert types(obs("50"), obs(None), rule) == [(PRICE_CHANGED, "50", None)]
    assert types(obs(None), obs(None), rule) == []


def test_database_upgrade_keeps_old_observations_and_allows_missing_price(tmp_path):
    path = tmp_path / "old.db"
    raw = sqlite3.connect(path)
    for statement in "".join(storage.MIGRATIONS[:4]).split(";"):
        if statement.strip():
            raw.execute(statement)
    raw.execute("PRAGMA user_version = 4")
    raw.execute("INSERT INTO sources VALUES ('s', 'http://x.test')")
    raw.execute("INSERT INTO collection_runs(source_id, started_at, status) VALUES ('s', 't', 'success')")
    raw.execute("INSERT INTO products VALUES ('s', '00123', 'N', 'http://x.test/1', 't')")
    raw.execute("INSERT INTO observations(run_id, source_id, product_key, name, url, price, currency, availability, observed_at) "
                "VALUES (1, 's', '00123', 'N', 'http://x.test/1', '9.90', 'UAH', 'IN_STOCK', 't')")
    raw.commit()
    raw.close()

    conn = storage.connect(path)
    assert conn.execute("PRAGMA user_version").fetchone()[0] == len(storage.MIGRATIONS)
    assert conn.execute("SELECT product_key, price FROM observations").fetchall()[0][:] == ("00123", "9.90")
    run = storage.start_run(conn, "s")
    storage.commit_snapshot(conn, run, "s", [obs(None)])
    assert conn.execute("SELECT price FROM observations WHERE run_id = ?", (run,)).fetchone()[0] is None


def test_report_shows_an_empty_price_cell(conn, tmp_path):
    storage.ensure_source(conn, "s", "http://x.test")
    run = storage.start_run(conn, "s")
    storage.commit_snapshot(conn, run, "s", [obs(None)])
    now = datetime.now(timezone.utc)
    path = reports.write_report(conn, tmp_path, "s", now - timedelta(days=1), now + timedelta(days=1))
    row = list(load_workbook(path)["Current"].iter_rows(min_row=2, values_only=True))[0]
    assert row[1] == "p1" and row[4] is None
