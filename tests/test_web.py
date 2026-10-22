import hashlib
import sqlite3
from decimal import Decimal as D
from io import BytesIO
from pathlib import Path

import pytest

pytest.importorskip("fastapi")
from fastapi.testclient import TestClient  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

from conftest import make_source  # noqa: E402
from price_monitor import demo_site, storage  # noqa: E402
from price_monitor.collector import collect_once  # noqa: E402
from price_monitor.models import Availability, ProductObservation  # noqa: E402
from price_monitor.web import auth  # noqa: E402
from price_monitor.web.app import create_app  # noqa: E402
from price_monitor.web.db import open_readonly  # noqa: E402

TOKEN = "test-token-0123456789"
BEARER = {"Authorization": f"Bearer {TOKEN}"}


@pytest.fixture
def db(tmp_path, demo):
    """A database filled by real collections of the demo catalog: baseline, price drop, stock changes."""
    url, state = demo
    path = tmp_path / "monitor.db"
    conn = storage.connect(path)
    source = make_source(url)
    collect_once(conn, source)
    storage.set_rule(conn, "demo", "00123", threshold=D("20"))
    for scenario in ("price_drop", "out_of_stock"):
        demo_site.write_scenario(state, scenario)
        collect_once(conn, source)
    conn.close()
    return path


def client_for(path: Path) -> TestClient:
    return TestClient(create_app(path, TOKEN, {"demo": make_source("http://x.test")}), follow_redirects=False)


@pytest.fixture
def client(db):
    return client_for(db)


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def test_the_dashboard_refuses_to_start_without_a_strong_token(tmp_path):
    for bad in (None, "", "short"):
        with pytest.raises(auth.TokenError):
            create_app(tmp_path / "x.db", bad)


def test_pages_need_the_token(client):
    page = client.get("/products?q=lamp", headers={"Accept": "text/html"})
    assert page.status_code == 303
    assert page.headers["location"] == "/login?next=%2Fproducts%3Fq%3Dlamp"
    assert client.get("/").status_code == 401
    assert client.get("/", headers={"Authorization": "Bearer wrong-token-0123456789"}).status_code == 401
    htmx = client.get("/products", headers={"HX-Request": "true"})
    assert htmx.status_code == 401 and htmx.headers["HX-Redirect"] == "/login"
    assert client.get("/report.xlsx?source=demo").status_code == 401
    # public: the sign-in page, the health check and the static files
    for path in ("/login", "/healthz", "/static/app.css", "/static/htmx.min.js"):
        assert client.get(path).status_code == 200


def test_sign_in_sets_an_http_only_cookie_and_only_local_redirects(client):
    assert client.post("/login", data={"token": "wrong"}).status_code == 401
    response = client.post("/login", data={"token": TOKEN, "next": "//evil.test/x"})
    assert response.status_code == 303 and response.headers["location"] == "/"
    cookie = response.headers["set-cookie"]
    assert "HttpOnly" in cookie and "samesite=strict" in cookie.lower()
    assert TOKEN not in cookie
    assert client.get("/").status_code == 200
    assert client.post("/login", data={"token": TOKEN, "next": "/events?kind=REMOVED"}).headers["location"] == (
        "/events?kind=REMOVED"
    )
    client.post("/logout")
    assert client.get("/").status_code == 401


def test_every_page_answers_and_sends_security_headers(client):
    for path in ("/", "/products", "/product?source=demo&key=00123", "/events", "/report", "/status"):
        response = client.get(path, headers=BEARER)
        assert response.status_code == 200, path
        assert "script-src 'self'" in response.headers["content-security-policy"]
        assert response.headers["x-frame-options"] == "DENY"
    assert client.get("/nothing", headers=BEARER).status_code == 404
    assert client.get("/products?source=nope", headers=BEARER).status_code == 404


def test_keys_are_strings_and_the_price_history_is_exact(client):
    page = client.get("/products", headers=BEARER).text
    assert "<code>00123</code>" in page and "<code>123</code>" in page

    lamp = client.get("/product", params={"source": "demo", "key": "00123"}, headers=BEARER).text
    assert "Desk lamp Nord" in lamp and "spare bulb" not in lamp
    assert '"price": "24.90"' in lamp and '"price": "19.90"' in lamp  # chart data keeps decimal strings
    assert "19.90 USD" in lamp
    assert "Ціна нижче порогу" in lamp

    bulb = client.get("/product", params={"source": "demo", "key": "123"}, headers=BEARER).text
    assert "spare bulb" in bulb and "19.90" not in bulb
    assert client.get("/product", params={"source": "demo", "key": "0123"}, headers=BEARER).status_code == 404


def test_search_and_filters(client):
    def keys(**params):
        text = client.get("/products", params=params, headers=BEARER).text
        return {k for k in ("00123", "123", "00789", "01001", "01006") if f"<code>{k}</code>" in text}

    assert keys(q="LAMP") == {"00123", "123"}  # case-insensitive
    assert keys(q="00123") == {"00123"}
    assert keys(availability="OUT_OF_STOCK") == {"01001"}
    assert keys(availability="UNKNOWN") == {"01006"}
    assert keys(tracked="1") == {"00123"}
    assert "Нічого не знайдено" in client.get("/products", params={"q": "zzz"}, headers=BEARER).text


def test_stock_directions_are_a_view_on_stored_events(client, db):
    def rows(kind):
        text = client.get("/events", params={"kind": kind}, headers=BEARER).text
        return {k for k in ("00123", "00789", "01001", "01003") if f"<code>{k}</code>" in text}

    assert rows("OUT_OF_STOCK") == {"01001"}
    assert rows("BACK_IN_STOCK") == {"00789"}
    assert rows("BELOW_THRESHOLD") == {"00123"}
    assert rows("PRICE_CHANGED") == {"00123", "01003"}
    stored = {r[0] for r in sqlite3.connect(db).execute("SELECT DISTINCT type FROM events")}
    assert "OUT_OF_STOCK" not in stored and "AVAILABILITY_CHANGED" in stored

    assert "Знайдено: 0" in client.get("/events", params={"since": "2000-01-01", "until": "2000-01-02"}, headers=BEARER).text
    bad = client.get("/events", params={"since": "yesterday"}, headers=BEARER)
    assert bad.status_code == 200 and "РРРР-ММ-ДД" in bad.text


def test_text_from_shop_pages_is_escaped_and_unsafe_links_are_dropped(tmp_path):
    path = tmp_path / "x.db"
    conn = storage.connect(path)
    storage.ensure_source(conn, "demo", "http://x.test")
    run_id = storage.start_run(conn, "demo")
    storage.commit_snapshot(conn, run_id, "demo", [
        ProductObservation("<b>1</b>", "<script>alert(1)</script>", "javascript:alert(1)", D("1"), "USD", Availability.IN_STOCK),
    ])
    conn.close()
    client = client_for(path)
    listing = client.get("/products", headers=BEARER).text
    assert "<script>alert(1)</script>" not in listing and "&lt;script&gt;alert(1)&lt;/script&gt;" in listing
    product = client.get("/product", params={"source": "demo", "key": "<b>1</b>"}, headers=BEARER).text
    assert "<b>1</b>" not in product and "javascript:" not in product


def test_the_dashboard_never_writes(client, db):
    before = digest(db)
    for path in ("/", "/products?q=lamp", "/product?source=demo&key=00123", "/events", "/status",
                 "/report.xlsx?source=demo"):
        assert client.get(path, headers=BEARER).status_code == 200
    assert digest(db) == before
    # a WAL reader may create the shared-memory index, but never writes a page into the log
    wal = db.with_name(db.name + "-wal")
    assert not wal.exists() or wal.stat().st_size == 0
    conn = open_readonly(db)
    with pytest.raises(sqlite3.OperationalError):
        conn.execute("DELETE FROM events")
    conn.close()


def test_status_hides_the_chat_id_and_bot_tokens(db):
    conn = storage.connect(db)
    conn.execute(
        "INSERT INTO notification_outbox(chat_id, dedupe_key, text, status, attempts, next_attempt_at, last_error, created_at) "
        "VALUES (424242, 'k', 't', 'failed', 5, '2026-01-01T00:00:00+00:00', "
        "'Client error for url https://api.telegram.org/bot123456:ABC-def_ghi/sendMessage', '2026-01-01T00:00:00+00:00')"
    )
    conn.commit()
    conn.close()
    html = client_for(db).get("/status", headers=BEARER).text
    assert "424242" not in html and "ABC-def_ghi" not in html
    assert "bot&lt;hidden&gt;" in html


def test_sources_switched_off_in_the_configuration_are_not_shown(db):
    conn = storage.connect(db)
    conn.execute("INSERT INTO sources(id, base_url) VALUES ('allo', 'https://allo.ua')")
    conn.commit()
    conn.close()
    html = client_for(db).get("/", headers=BEARER).text
    assert "allo.ua" not in html and "demo" in html


def test_missing_or_foreign_database_is_explained_not_created(tmp_path):
    missing = tmp_path / "none.db"
    response = client_for(missing).get("/", headers=BEARER)
    assert response.status_code == 503 and "першого збору" in response.text
    assert not missing.exists()

    old = tmp_path / "old.db"
    sqlite3.connect(old).execute("PRAGMA user_version = 3")
    assert "старіша" in client_for(old).get("/status", headers=BEARER).text


def test_empty_database_shows_empty_states(tmp_path):
    path = tmp_path / "empty.db"
    storage.connect(path).close()
    client = client_for(path)
    for page in ("/", "/products", "/report"):
        assert "ще немає жодного джерела" in client.get(page, headers=BEARER).text
    assert "Зборів ще не було" in client.get("/status", headers=BEARER).text


def test_xlsx_download_matches_the_database_and_cleans_up(client, tmp_path, monkeypatch):
    import tempfile

    made = []
    real_mkdtemp = tempfile.mkdtemp

    def mkdtemp(**kwargs):
        made.append(Path(real_mkdtemp(dir=tmp_path, **kwargs)))
        return str(made[-1])

    monkeypatch.setattr(tempfile, "mkdtemp", mkdtemp)
    response = client.get("/report.xlsx", params={"source": "demo", "since": "2000-01-01", "until": "2100-01-01"}, headers=BEARER)
    assert response.status_code == 200
    assert response.headers["content-disposition"].startswith('attachment; filename="report-demo-')
    wb = load_workbook(BytesIO(response.content))
    current = {row[1]: row for row in wb["Current"].iter_rows(min_row=2, values_only=True)}
    assert "00123" in current and "123" in current
    assert current["00123"][4] == pytest.approx(24.90)  # the out_of_stock scenario restores base prices
    history = {(r[1], r[8]) for r in wb["History"].iter_rows(min_row=2, values_only=True)}
    assert ("01001", "AVAILABILITY_CHANGED") in history
    assert made and not made[0].exists()  # the temporary file is gone after the download

    bad = client.get("/report.xlsx", params={"source": "demo", "since": "2026-02-01", "until": "2026-01-01"}, headers=BEARER)
    assert bad.status_code == 400 and "не пізніше" in bad.text
    assert client.get("/report.xlsx", params={"source": "nope"}, headers=BEARER).status_code == 404
